"""Kiểm tra deck Buổi 1: tiêu đề khớp outline, chữ trên slide có trong note, số từ phần "Nói"."""
import re, sys, json
from pptx import Presentation

deck, outline = sys.argv[1], sys.argv[2]
titles = {int(n): t.strip() for n, t in re.findall(r"\*\*Slide (\d+) — (.+?)\*\*", open(outline, encoding="utf-8").read())}
titles[40] = "Tài liệu tham khảo"

def norm(t):
    t = t.replace("&", " và ").lower()
    t = re.sub(r"[,.:;!?“”\"'‘’()\[\]·→–—~+−\-…=/•]", " ", t)
    return re.sub(r"\s+", " ", t).strip()

def say_words(note):
    say = note.split("\n\nHỏi lớp:")[0]
    say = re.sub(r"\n\(GV:.*", "", say, flags=re.S).replace("Nói:", "")
    return len(say.split())

prs = Presentation(deck)
rows = []
for i, sl in enumerate(prs.slides, 1):
    tshape = next((sh for sh in sl.shapes if sh.is_placeholder and sh.placeholder_format.type in (1, 3)), None)
    title = tshape.text_frame.text if tshape else ""
    note = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
    nn = norm(note)
    missing = []
    for sh in sl.shapes:
        if not sh.has_text_frame or sh is tshape or (sh.is_placeholder and sh.shape_id == tshape.shape_id):
            continue
        for line in sh.text_frame.text.split("\n"):
            t = norm(line)
            if len(t) < 3 or re.fullmatch(r"[\d %']+", t) or t in ("vs",):
                continue
            if t not in nn:
                missing.append(line.strip())
    rows.append({"n": i, "title_ok": title == titles.get(i), "title": title, "missing": missing,
                 "say_words": say_words(note), "parts_ok": all(k in note for k in ("Nói:", "Hỏi lớp:", "Chuyển ý:"))})
json.dump(rows, open(deck.replace(".pptx", ".check.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
for r in rows:
    flag = [] if r["title_ok"] else ["TITLE"]
    if r["missing"]: flag.append("MISSING " + " | ".join(r["missing"]))
    if not (60 <= r["say_words"] <= 150): flag.append(f"WORDS {r['say_words']}")
    if not r["parts_ok"]: flag.append("PARTS")
    if flag: print(r["n"], "; ".join(flag))
print("slides:", len(rows))
